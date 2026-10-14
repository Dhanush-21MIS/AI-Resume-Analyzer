import io

import pytest
from docx import Document

from app.document_parser import extract_resume_text


def create_docx_bytes() -> bytes:
    """Create a valid DOCX document in memory."""
    document = Document()

    document.add_heading("Resume", level=1)
    document.add_paragraph(
        "Python developer with experience in FastAPI and SQL."
    )

    buffer = io.BytesIO()
    document.save(buffer)

    return buffer.getvalue()


def create_pdf_bytes() -> bytes:
    """Create a valid, blank PDF document in memory."""
    import fitz

    pdf_document = fitz.open()
    pdf_document.new_page(width=612, height=792)

    file_bytes = pdf_document.tobytes()
    pdf_document.close()

    return file_bytes


def test_extract_text_from_docx():
    """Verify text extraction from a valid DOCX document."""
    file_bytes = create_docx_bytes()

    extracted_text = extract_resume_text(
        file_bytes,
        "resume.docx",
    )

    assert extracted_text
    assert "Python developer" in extracted_text
    assert "FastAPI" in extracted_text
    assert "SQL" in extracted_text


def test_pdf_without_readable_text_raises_value_error():
    """Verify that a blank PDF raises a clear validation error."""
    file_bytes = create_pdf_bytes()

    with pytest.raises(
        ValueError,
        match="No readable text was found",
    ):
        extract_resume_text(
            file_bytes,
            "resume.pdf",
        )


def test_corrupted_docx_is_handled():
    """Verify that corrupted DOCX data is not silently accepted."""
    corrupted_bytes = b"This is not a valid DOCX document."

    with pytest.raises(Exception):
        extract_resume_text(
            corrupted_bytes,
            "corrupted.docx",
        )