from io import BytesIO

import fitz
from docx import Document


def extract_resume_text(file_bytes: bytes, filename: str) -> str:
    """
    Extract readable text from a PDF or DOCX resume.

    Args:
        file_bytes: Uploaded document contents.
        filename: Original uploaded filename.

    Returns:
        Extracted resume text.

    Raises:
        ValueError: If the format is unsupported or no text is found.
    """

    extension = filename.lower().rsplit(".", 1)[-1]

    if extension == "pdf":
        text_parts = []

        with fitz.open(
            stream=file_bytes,
            filetype="pdf"
        ) as pdf_document:

            for page in pdf_document:
                page_text = page.get_text("text")

                if page_text.strip():
                    text_parts.append(page_text)

        text = "\n".join(text_parts)

    elif extension == "docx":
        document = Document(BytesIO(file_bytes))

        text_parts = []

        # Extract paragraphs.
        for paragraph in document.paragraphs:
            paragraph_text = paragraph.text.strip()

            if paragraph_text:
                text_parts.append(paragraph_text)

        # Extract text from tables.
        for table in document.tables:
            for row in table.rows:
                row_text = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]

                if row_text:
                    text_parts.append(" | ".join(row_text))

        text = "\n".join(text_parts)

    else:
        raise ValueError(
            "Unsupported file format. Upload a PDF or DOCX file."
        )

    text = text.strip()

    if not text:
        raise ValueError(
            "No readable text was found in the document. "
            "The PDF may contain scanned images and require OCR."
        )

    return text